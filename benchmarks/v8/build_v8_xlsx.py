#!/usr/bin/env python3
"""Build results_v8.xlsx from the v8 campaign (brain05, 30 Sep - 2 Oct 2026).

Inputs (same directory): plan_v8.csv, results_v8.csv, results_v8_w0A.csv
(reruns of the jobs that first ran with the withdrawn W = 0 rule; they
replace the original rows), env_v8.txt.  Rebuild any time; unfinished jobs
show as running / not reached.
"""
import csv, os, re, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "results_v8.xlsx")

plan = {r["id"]: r for r in csv.DictReader(open(os.path.join(HERE, "plan_v8.csv")))}
res = {r["id"]: r for r in csv.DictReader(open(os.path.join(HERE, "results_v8.csv")))}
rerun = {}
p = os.path.join(HERE, "results_v8_w0A.csv")
if os.path.exists(p):
    for r in csv.DictReader(open(p)):
        rerun[r["id"][:-2]] = r            # "brp-007-A" -> "brp-007"
for k, r in rerun.items():
    res[k] = dict(r, rerun="rule A rerun")

EXAMPLE_NAME = {
    "lease_timed": "Lease (TTL + fencing)", "swim_timed": "SWIM failure detector",
    "sensor_network_timed": "Sensor network (SCP 2014)", "brp_timed": "BRP (bounded retransmission)",
    "par_timed": "PAR (Bosnacki-Dams)", "alternating_bit_timed": "Alternating bit (ABP)",
    "chandra_toueg_timed": "Chandra-Toueg failure detector", "nrp_failure_detector_timed": "NRP failure detector",
    "nisshi_txn_timed": "nisshi transactions (KAFKA-17754)", "spanning_tree_timed": "Spanning tree (Perlman)",
    "yarn_scheduler_timed": "YARN scheduler", "retail_saga_timed": "Retail saga",
    "raft_leader_election": "Raft leader election", "three_pc_timed": "Three-phase commit",
    "two_pc_timed": "Two-phase commit", "comm_closed_leader_election": "Comm-closed leader election",
}
PAPER = {"lease_timed", "swim_timed", "sensor_network_timed", "brp_timed", "par_timed", "alternating_bit_timed",
         "chandra_toueg_timed", "nrp_failure_detector_timed", "nisshi_txn_timed", "spanning_tree_timed"}


def num(pat, s):
    m = re.findall(pat, s or "")
    return int(m[-1]) if m else None


def violations(example, r):
    """The example's own count of violating endings (None if it prints none)."""
    s = r.get("summary", "")
    if example == "nisshi_txn_timed":
        return num(r"dirty_reads=(\d+)", s)            # violations= double counts (stray commits are dirty too)
    if example in ("par_timed", "alternating_bit_timed"):
        a, b = num(r"p1_fails=(\d+)", s), num(r"p2_fails=(\d+)", s)
        return None if a is None and b is None else (a or 0) + (b or 0)
    if example == "brp_timed":
        vals = [num(rf"{k}=(\d+)", s) for k in ("p2_runs", "p4_runs")]
        return None if all(v is None for v in vals) else sum(v or 0 for v in vals)
    if example == "sensor_network_timed":
        a, b = num(r"sensor_failures=(\d+)", s), num(r"\bdead=(\d+)", s)
        return None if a is None and b is None else (a or 0) + (b or 0)
    if example == "retail_saga_timed":
        a, b = num(r"oversold=(\d+)", s), num(r"\blate=(\d+)", s)
        return None if a is None and b is None else (a or 0) + (b or 0)
    v = r.get("violations", "")
    return int(v) if v not in ("", None) else num(r"violations=(\d+)", s)


def state(i):
    """(status, explored, wall, violations, impossible, note) of plan row i."""
    p = plan[i]
    r = res.get(i)
    if r is None:
        return ("running" if os.path.exists(os.path.join(HERE, "status_v8.json")) else "not run", None, None, None, None, "")
    st = r["status"]
    note = r.get("rerun", "")
    if st.startswith("DNF"):
        return ("DNF", None, float(r["wall_s"] or 0), None, None, (r.get("progress") or "") + ("; " + note if note else ""))
    if st.startswith("error"):
        return ("error", None, float(r["wall_s"] or 0), None, None, "killed: ran with the withdrawn W = 0 rule; see rerun" if i in rerun else st)
    if st.startswith("not run"):
        return ("not run", None, None, None, None, "")
    ex, bl = r.get("execs"), r.get("blocked")
    explored = int(ex) + int(bl) if ex not in ("", None) and bl not in ("", None) else None
    fire = "assertion failed" in st
    v = violations(p["example"], r)
    imp = r.get("timeline_impossible")
    return ("FIRE (first violation)" if fire else "finished", explored, float(r["wall_s"] or 0), v,
            int(imp) if imp not in ("", None) else None, note)


# ---- pair the timed and baseline sides of each cell ----------------------------------------
pairs = collections.OrderedDict()
for i, p in plan.items():
    key = (p["group"], p["example"], p["cell"])
    pairs.setdefault(key, {})[p["side"]] = i

wb = Workbook()
H = Font(bold=True, color="FFFFFF")
HF = PatternFill("solid", fgColor="305496")
GOOD = PatternFill("solid", fgColor="E2EFDA")
BAD = PatternFill("solid", fgColor="FCE4D6")
GREY = PatternFill("solid", fgColor="EDEDED")


def header(ws, cols, widths):
    ws.append(cols)
    for c in ws[1]:
        c.font, c.fill = H, HF
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for k, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(k)].width = w
    ws.freeze_panes = "A2"


def ratio(a, b):
    return round(a / b, 2) if a and b else None


# ---- Pairs (every cell) -----------------------------------------------------------------------
ws = wb.active
ws.title = "All cells"
cols = ["Benchmark", "In paper set", "Cell", "Purpose", "Timed status", "Timed explored (execs+blocked)", "Timed wall (s)",
        "Timed violations", "Baseline status", "Baseline explored", "Baseline wall (s)", "Baseline violations",
        "Explored ratio (baseline / timed)", "Wall ratio (baseline / timed)", "Cores", "Timeline-impossible (timed)",
        "Notes", "Timed args"]
header(ws, cols, [26, 8, 46, 10, 22, 16, 12, 12, 22, 16, 12, 12, 12, 12, 7, 11, 44, 60])
for (g, ex, cell), sides in pairs.items():
    t = state(sides["timed"]) if "timed" in sides else (None,) * 6
    b = state(sides["baseline"]) if "baseline" in sides else (None,) * 6
    any_id = sides.get("timed") or sides.get("baseline")
    notes = "; ".join(x for x in (t[5], b[5]) if x)
    row = [EXAMPLE_NAME.get(ex, ex), "yes" if ex in PAPER else "", cell, plan[any_id]["purpose"],
           t[0], t[1], t[2], t[3], b[0], b[1], b[2], b[3], ratio(b[1], t[1]), ratio(b[2], t[2]),
           plan[any_id]["cpus"], t[4], notes, plan[sides["timed"]]["args"] if "timed" in sides else ""]
    ws.append(row)
    rr = ws.max_row
    if t[3] == 0 and (b[3] or 0) > 0:
        ws.cell(rr, 8).fill = GOOD
        ws.cell(rr, 12).fill = BAD
    for k in (6, 10):
        ws.cell(rr, k).number_format = "#,##0"
    for k in (7, 11):
        ws.cell(rr, k).number_format = "#,##0.0"

# ---- Headlines ---------------------------------------------------------------------------------
hs = wb.create_sheet("Headline cells", 0)
header(hs, ["Benchmark", "Cell", "Timed explored", "Baseline explored", "Explored ratio", "Timed wall (s)",
            "Baseline wall (s)", "Wall ratio", "Timed violations", "Baseline violations", "Reading"],
       [26, 50, 15, 15, 9, 11, 11, 9, 11, 12, 70])
for (g, ex, cell), sides in pairs.items():
    any_id = sides.get("timed") or sides.get("baseline")
    if plan[any_id]["purpose"] != "headline" or ex not in PAPER:
        continue
    t = state(sides["timed"]) if "timed" in sides else (None,) * 6
    b = state(sides["baseline"]) if "baseline" in sides else (None,) * 6
    if t[0] is None or b[0] is None:
        continue
    reading = []
    if t[3] == 0 and (b[3] or 0) > 0:
        reading.append(f"precision: untimed reports {b[3]:,} spurious violating endings, timed none")
    if t[1] and b[1]:
        r = b[1] / t[1]
        reading.append(f"timed explores {r:.1f}x fewer" if r > 1.05 else (f"timed explores {1/r:.1f}x more" if r < 0.95 else "same exploration size"))
    hs.append([EXAMPLE_NAME.get(ex, ex), cell, t[1], b[1], ratio(b[1], t[1]), t[2], b[2], ratio(b[2], t[2]), t[3], b[3],
               "; ".join(reading)])
    rr = hs.max_row
    for k in (3, 4):
        hs.cell(rr, k).number_format = "#,##0"

# ---- Ceiling grids for lease and SWIM -------------------------------------------------------------
def grid(title, group, rx_size):
    gs = wb.create_sheet(title)
    cells = collections.defaultdict(dict)
    fams, sizes = set(), set()
    for (g, ex, cell), sides in pairs.items():
        if g != group:
            continue
        m = re.search(rx_size, cell)
        if not m:
            continue
        size = m.group(0)
        fam = re.sub(rx_size, "", cell).replace("  ", " ").strip()
        fams.add(fam); sizes.add(size)
        for side, i in sides.items():
            s = state(i)
            txt = s[0] if s[2] is None else f"{s[0]} {s[2]:,.0f} s"
            if s[1]:
                txt += f", {s[1]:,} explored"
            if s[3] is not None:
                txt += f", viol {s[3]:,}"
            if s[0] == "DNF" and s[5]:
                txt += f" ({s[5]})"
            cells[(fam, size)][side] = txt
    sizes = sorted(sizes, key=lambda z: [int(x) for x in re.findall(r"\d+", z)])
    header(gs, ["Family / side"] + sizes, [48] + [30] * len(sizes))
    for fam in sorted(fams):
        for side in ("timed", "baseline"):
            row = [f"{fam} [{side}]"] + [cells[(fam, z)].get(side, "") for z in sizes]
            gs.append(row)
            for c in gs[gs.max_row][1:]:
                c.alignment = Alignment(wrap_text=True, vertical="top")
                if c.value and c.value.startswith("DNF"):
                    c.fill = GREY


grid("Lease ceiling", "lease", r"C=\d+ R=\d+")
grid("SWIM ceiling", "swim", r"N=\d+ R=\d+")

# ---- Method --------------------------------------------------------------------------------------
ms = wb.create_sheet("Method")
env = open(os.path.join(HERE, "env_v8.txt")).read().splitlines() if os.path.exists(os.path.join(HERE, "env_v8.txt")) else []
lines = [
    "v8 campaign, brain05 (256 cores, 2 TB RAM), 30 Sep 2026 13:44 to 2 Oct 2026 08:00 (hard deadline).",
    "Algorithm: the new exploration (every send may be dropped; dropped is the canonical send outcome; consistency = the full timed constraint system at every step).",
    "Sides: 'timed' = with_timed; 'baseline' = the same program without timing (untimed Must, every send lossy too). Each side is a separate job, pinned to its own cores.",
    "Explored = execs + blocked endings (both are behaviours checked). Timeline-impossible endings must be 0 under the new algorithm (all finished jobs: 0).",
    "Violations = the example's own count of violating endings under --keep-going (nisshi: dirty_reads; PAR/ABP: p1+p2 fails; BRP: p2+p4 runs; sensor: false failures + false deaths).",
    "FIRE (first violation) = run without --keep-going that stopped at the first certified violation.",
    "DNF = still running at the deadline and killed; 'progress' gives how far it got (explored so far).",
    "W = 0 rule: the campaign started 13:44 with a W = 0 timeout rule that was withdrawn at 16:40 (rule A: a message arriving exactly at the deadline may be read or missed for every W). The 18 jobs that ran a W = 0 receive under the withdrawn rule (timed BRP with property 4/all, timed raft) were rerun with rule A; their rerun rows replace the originals here (Notes: 'rule A rerun'). brp-007/022 reruns used 4 cores instead of 6.",
    "Calibration findings (30 Sep): SWIM boundary is FIRE iff Ws <= 2U + sd - min(L, max(0, 2L - Wp)) for N >= 3 (sd once); Chandra-Toueg: timed explores about 2x more than baseline (precision-only); NRP: runs dominated by blocked endings (loss + eviction guard); lease rounds >= 2 needed the RoundStart/Grant tag fix.",
    "",
    "Build environment:",
] + env
for l in lines:
    ms.append([l])
ms.column_dimensions["A"].width = 160

wb.save(OUT)
print("wrote", OUT, "| cells:", len(pairs), "| result rows:", len(res), "| reruns applied:", len(rerun))
