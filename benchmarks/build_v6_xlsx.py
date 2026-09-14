#!/usr/bin/env python3
"""Build results_v6.xlsx from results_v6_grid.csv and results_v6_lease_swim.csv.

Re-runnable: fetch the CSVs from brain05 (~/traceforge-v6/benchmarks/) and rerun
after more grid cells finish.

Sheets:
  Results                 one row per grid cell (3PC, plain LE, inbox LE), timed
                          and untimed side by side, pruning and wall-time ratio;
  Lease & SWIM comparison the --keep-going cells: untimed and timed execution
                          counts side by side, with the number of violating
                          executions the untimed checker reports;
  Lease & SWIM boundaries predicted verdict (from the published inequality)
                          versus the measured one, both sides;
  Legend & method, Cell guide.
"""
import csv, os, re, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))

def load(name):
    p = os.path.join(HERE, name)
    return list(csv.DictReader(open(p))) if os.path.exists(p) else []

grid = load("results_v6_grid.csv")
ls   = load("results_v6_lease_swim.csv")
# Runs interrupted by the brain05 reboot of 2026-09-10 are recorded as DNF rows
# in a separate file (kept out of results_v6_grid.csv so the sweep's resume
# logic still re-runs them). A finished row for the same (cell, side) wins.
interrupted = load("results_v6_grid_interrupted.csv")
have = {(r["protocol"], r["N"], r["rounds"], r["cell"], r["side"]) for r in grid}
grid += [r for r in interrupted if (r["protocol"], r["N"], r["rounds"], r["cell"], r["side"]) not in have]

def num(v):
    try: return int(v)
    except (TypeError, ValueError):
        try: return float(v)
        except (TypeError, ValueError): return v or ""

def params(cmd):
    """(L, U, sd, W) in ticks from a grid command line: LE passes absolute
    ticks (--l/--u/--sd/--w), 3PC passes ratios of U."""
    def flag(name):
        m = re.search(r"--%s\s+([0-9.]+)" % re.escape(name), cmd)
        return float(m.group(1)) if m else None
    u = flag("u")
    if u is None:
        return ("", "", "", "")
    if flag("w") is not None:
        l, sd, w = flag("l"), flag("sd"), flag("w")
    else:
        lr, sr, wr = flag("l-ratio"), flag("sd-ratio"), flag("w-ratio")
        l  = round(lr * u) if lr is not None else None
        sd = round(sr * u) if sr is not None else None
        w  = round(wr * u) if wr is not None else None
    fmt = lambda v: "" if v is None else int(v)
    return (fmt(l), fmt(u), fmt(sd), fmt(w))

HDR   = Font(bold=True, color="FFFFFF")
HFILL = PatternFill("solid", fgColor="2C5F8A")
WARN  = PatternFill("solid", fgColor="FCE4D6")
GOOD  = PatternFill("solid", fgColor="E2EFDA")
PEND  = PatternFill("solid", fgColor="FFF2CC")
BOLD  = Font(bold=True)
THIN  = Border(*[Side(style="thin", color="BFBFBF")] * 4)

def style_header(ws, row, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.font, cell.fill, cell.border = HDR, HFILL, THIN
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

def autosize(ws, maxw=44):
    for col in ws.columns:
        w = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(w + 2, maxw)

def border_row(ws, i, ncol):
    for c in range(1, ncol + 1):
        ws.cell(row=i, column=c).border = THIN

def pair_up(rows):
    """Group per-side rows into {cell_key: {"timed": row, "baseline": row}}."""
    cells = {}
    for r in rows:
        k = (r["protocol"], int(r["N"]), int(r["rounds"]), r["cell"])
        cells.setdefault(k, {})[r["side"]] = r
    return cells

def side(c, s, field):
    r = c.get(s)
    return num(r[field]) if r and r.get(field) not in (None, "") else ""

def verdict(c, s):
    r = c.get(s)
    return r["verdict"] if r else "running"

wb = Workbook()

# ---------------------------------------------------------------- Results ---
ws = wb.active; ws.title = "Results"
ws["A1"] = "TraceForge v6 sweep: timed vs untimed baseline (exactly-k timed inbox, no time caps)"
ws["A1"].font = Font(bold=True, size=14)
ws["A2"] = ("All runs on brain05 (256 cores). Timed and untimed are separate invocations with no time cap; "
            "a cell still running when this file was built shows 'running' (yellow); a run that was killed by the "
            "brain05 reboot of 2026-09-10 shows 'DNF (rebooted)' with the wall time it had run until the reboot (orange). "
            "Execution counts are machine-independent and are the primary metric; wall times are indicative "
            "(shared work-queue pool, 64 workers per run, three or four runs sharing the machine).")
ws["A3"] = f"Generated {datetime.date.today().isoformat()} from results_v6_grid.csv"

cols = ["Protocol", "N", "Rounds/Ballots", "Cell",
        "Timed execs", "Timed blocked", "Timed verdict", "Timed wall (s)", "Timed CPU (s)",
        "Untimed execs", "Untimed blocked", "Untimed verdict", "Untimed wall (s)", "Untimed CPU (s)",
        "Pruning %", "Wall ratio (untimed / timed)",
        "L", "U", "sd", "W", "Command"]
ws.append([]); ws.append(cols); style_header(ws, 5, len(cols))

gcells = pair_up(grid)
# planned cells with no finished side yet appear as fully "running" rows
try:
    import sweep_v6
    for pc in sweep_v6.grid_cells():
        k = (pc["protocol"], int(pc["N"]), int(pc["rounds"]), pc["cell"])
        if k not in gcells:
            gcells[k] = {"_planned_cmd": " ".join(pc["cmd"]).replace(sweep_v6.BIN + os.sep, "")}
except Exception as e:  # the sheet still builds from the CSV alone
    print("note: could not load planned cells:", e)
pending = 0
for k in sorted(gcells):
    c = gcells[k]
    cmd = (c.get("timed") or c.get("baseline") or {"cmd": c.get("_planned_cmd", "")})["cmd"]
    L, U, SD, W = params(cmd)
    te, be = side(c, "timed", "execs"), side(c, "baseline", "execs")
    tw, bw = side(c, "timed", "wall_s"), side(c, "baseline", "wall_s")
    prune = round(100.0 * (be - te) / be, 1) if isinstance(te, int) and isinstance(be, int) and be else ""
    ratio = round(bw / tw, 2) if isinstance(tw, (int, float)) and isinstance(bw, (int, float)) and tw else ""
    ws.append([k[0], k[1], k[2], k[3],
               te, side(c, "timed", "block"), verdict(c, "timed"), tw, side(c, "timed", "cpu_s"),
               be, side(c, "baseline", "block"), verdict(c, "baseline"), bw, side(c, "baseline", "cpu_s"),
               prune, ratio, L, U, SD, W, cmd.replace(" --mode timed", "").replace(" --mode baseline", "")])
    i = ws.max_row; border_row(ws, i, len(cols))
    for col in (5, 6, 8, 9, 10, 11, 13, 14):
        ws.cell(row=i, column=col).number_format = "#,##0"
    ws.cell(row=i, column=15).number_format = '0.0"%"'
    vt, vb = verdict(c, "timed"), verdict(c, "baseline")
    if "running" in (vt, vb):
        pending += 1
        for col in range(1, len(cols) + 1):
            ws.cell(row=i, column=col).fill = PEND
    elif prune != "" and prune >= 99:
        ws.cell(row=i, column=15).fill = GOOD
        ws.cell(row=i, column=15).font = BOLD
    if any(v.startswith("DNF") for v in (vt, vb)):
        for col in range(1, len(cols) + 1):
            ws.cell(row=i, column=col).fill = WARN
    elif vt not in ("hold", "running") or vb not in ("hold", "running"):
        for col in range(1, len(cols) + 1):
            ws.cell(row=i, column=col).fill = WARN
if pending:
    ws.append([])
    ws.append([f"NOTE: {pending} cells have at least one side still running on brain05 (yellow rows). "
               "Rerun build_v6_xlsx.py after fetching newer CSVs."])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True, color="9C5700")
    ws.cell(row=ws.max_row, column=1).fill = PEND
ws.freeze_panes = "E6"; autosize(ws)

# ---------------------------------------------- Lease & SWIM comparison ----
ws2 = wb.create_sheet("Lease & SWIM comparison")
ws2["A1"] = "Lease and SWIM: untimed vs timed execution counts (--keep-going cells)"
ws2["A1"].font = Font(bold=True, size=14)
ws2["A2"] = ("Both sides explore their whole state space (the checker keeps going after a violation instead of aborting). "
             "'Violations' is how many explored executions break the property: for the untimed checker every one of them "
             "is a spurious counterexample (it cannot tell a short pause or a fast refutation from a long one); the timed "
             "checker reports 0 and proves the property at these parameters. Total = execs + blocked. "
             "Lease cells sit one tick below the split-brain boundary (unfenced) or at the boundary with fencing on; "
             "SWIM cells sit one tick above the false-positive boundary.")
ws2["A2"].alignment = Alignment(wrap_text=True)
ws2.row_dimensions[2].height = 60
c2 = ["Protocol", "N (clients / members)", "Rounds", "Cell (parameters)",
      "Timed execs", "Timed blocked", "Timed total", "Timed violations", "Timed verdict", "Timed wall (s)", "Timed CPU (s)",
      "Untimed execs", "Untimed blocked", "Untimed total", "Untimed violations", "Untimed verdict", "Untimed wall (s)", "Untimed CPU (s)",
      "Total ratio (untimed / timed)"]
ws2.append([]); ws2.append(c2); style_header(ws2, 5, len(c2))
cmp_cells = pair_up([r for r in ls if r["variant"] == "cmp"])
for k in sorted(cmp_cells):
    c = cmp_cells[k]
    te, tb = side(c, "timed", "execs"), side(c, "timed", "block")
    be, bb = side(c, "baseline", "execs"), side(c, "baseline", "block")
    tt = te + tb if isinstance(te, int) and isinstance(tb, int) else ""
    bt = be + bb if isinstance(be, int) and isinstance(bb, int) else ""
    ratio = round(bt / tt, 2) if isinstance(tt, int) and isinstance(bt, int) and tt else ""
    ws2.append([k[0], k[1], k[2], k[3],
                te, tb, tt, side(c, "timed", "violations"), verdict(c, "timed"),
                side(c, "timed", "wall_s"), side(c, "timed", "cpu_s"),
                be, bb, bt, side(c, "baseline", "violations"), verdict(c, "baseline"),
                side(c, "baseline", "wall_s"), side(c, "baseline", "cpu_s"), ratio])
    i = ws2.max_row; border_row(ws2, i, len(c2))
    for col in (5, 6, 7, 8, 10, 11, 12, 13, 14, 15, 17, 18):
        ws2.cell(row=i, column=col).number_format = "#,##0"
    ok = (verdict(c, "timed") == c["timed"]["expected"] if "timed" in c else True) and \
         (verdict(c, "baseline") == c["baseline"]["expected"] if "baseline" in c else True)
    if not ok:
        for col in range(1, len(c2) + 1):
            ws2.cell(row=i, column=col).fill = WARN
ws2.freeze_panes = "E6"; autosize(ws2)

# ------------------------------------------------- Lease & SWIM boundaries ---
ws3 = wb.create_sheet("Lease & SWIM boundaries")
ws3["A1"] = "Boundary validation: predicted verdict vs measured (abort-on-first-violation runs)"
ws3["A1"].font = Font(bold=True, size=14)
ws3["A2"] = "Lease:  FIRE  iff  pause >= TTL - 2(U - L + sd)      (checked at R = 1, 2, 3)"
ws3["A3"] = "SWIM:   FIRE  iff  W_suspect <= 2U + 2sd - min(L, max(0, 2L - W_probe))      (checked at R = 1, 2, 3)"
bnd = [r for r in ls if r["variant"] == "boundary"]
agree = sum(1 for r in bnd if r["verdict"] == r["expected"])
ws3["A4"] = f"{agree} / {len(bnd)} rows (timed and untimed sides) match the predicted verdict"
ws3["A4"].font = Font(bold=True, color=("2E6B3E" if agree == len(bnd) and bnd else "C00000"))
c3 = ["Protocol", "Rounds", "Cell (parameters)", "Side", "Predicted", "Measured", "Match?",
      "Execs", "Blocked", "Wall (s)"]
ws3.append([]); ws3.append(c3); style_header(ws3, 6, len(c3))
for r in sorted(bnd, key=lambda r: (r["protocol"], int(r["rounds"]), r["cell"], r["side"])):
    ok = r["verdict"] == r["expected"]
    ws3.append([r["protocol"], num(r["rounds"]), r["cell"], "timed" if r["side"] == "timed" else "untimed",
                r["expected"], r["verdict"], "yes" if ok else "NO",
                num(r["execs"]), num(r["block"]), num(r["wall_s"])])
    i = ws3.max_row; border_row(ws3, i, len(c3))
    ws3.cell(row=i, column=7).fill = GOOD if ok else WARN
    if not ok:
        ws3.cell(row=i, column=7).font = Font(bold=True, color="C00000")
ws3.freeze_panes = "D7"; autosize(ws3)

# ------------------------------------------------------- Legend & method ----
ws4 = wb.create_sheet("Legend & method")
for line in [
    ("Method (v6, September 2026)", True),
    ("Every cell runs the identical program twice: once under the untimed baseline, once timed. Only the timing configuration differs.", False),
    ("Timed and untimed are separate process invocations with NO time cap; the largest cells run for days and a side that has not finished is marked 'running'.", False),
    ("The timed quorum inbox collects exactly k = floor(N/2)+1 acknowledgements within W (the range form used before September 2026 was retired), so the untimed inbox counts changed versus the v5 sheet: e.g. inbox N=3 B=1 untimed 680 (was 1,290), N=5 B=1 untimed 31,694,957 (was 278,558,488). Timed counts at N=3 and N=4 are unchanged.", False),
    ("Exploration uses the shared work-queue pool with 64 workers per run (8 for lease/SWIM); counts are identical to single-threaded exploration (verified on every protocol) and the pool's completion-gate parity fix of 2026-09-05 is included.", False),
    ("Execution counts are machine-independent and are the primary metric; wall and CPU times are from brain05 with several runs sharing the machine, so treat them as indicative. CPU time is the fairer cost measure across cells.", False),
    ("", False),
    ("Columns", True),
    ("Timed/Untimed execs", False), ("  Number of explored complete executions; blocked = executions that ended blocked (an unmatched receive, or a kept-going assertion violation).", False),
    ("Pruning %", False), ("  100 * (untimed - timed) / untimed. Blank when either side is still running.", False),
    ("Wall ratio", False), ("  untimed wall / timed wall for the same cell (both with the same worker count).", False),
    ("Verdict", False), ("  hold = no assertion fired; FIRE = a certified counterexample (or, in keep-going runs, violations > 0); running = not finished when the file was built.", False),
    ("", False),
    ("Regimes (U = 20 ticks)", True),
    ("A  tight intra-DC:   L=18, sd=2,  W=200   (L/U=0.9,  sd/U=0.1,  W/U=10)", False),
    ("B  intra-DC TCP:     L=6,  sd=0,  W=60    (L/U=0.3,  sd/U=0,    W/U=3)", False),
    ("C  inter-AZ:         L=2,  sd=5,  W=100   (L/U=0.1,  sd/U=0.25, W/U=5)", False),
    ("F  adversarial:      L=0,  sd=10, W=40    (L/U=0,    sd/U=0.5,  W/U=2)  <- widest message window; the completeness control (timed must equal untimed)", False),
    ("", False),
    ("Lease & SWIM comparison sheet", True),
    ("Cells run with --keep-going: the checker records a violation and continues instead of aborting, so the untimed side explores its whole space and reports how many executions violate the property. Lease: violations = stale writes applied; SWIM: violations = dead declarations of a live member (false positives).", False),
    ("Lease cells: unfenced at pause = boundary - 1 (timed proves safety, untimed cannot) and fenced at pause = boundary (both hold). SWIM cells: W_suspect = boundary + 1.", False),
    ("", False),
    ("Lease & SWIM boundaries sheet", True),
    ("Each cell carries the verdict predicted by the published inequality; 'Match?' compares it to what the tool returned, for the timed and the untimed side. The untimed side FIREs on every unfenced lease and every SWIM cell by construction.", False),
]:
    ws4.append([line[0]])
    if line[1]:
        ws4.cell(row=ws4.max_row, column=1).font = Font(bold=True, size=12)
ws4.column_dimensions["A"].width = 130

# --------------------------------------------------- Cell column guide -----
ws5 = wb.create_sheet("Cell guide")
ws5["A1"] = "What each value in the 'Cell' column means"; ws5["A1"].font = Font(bold=True, size=14)
ws5["A2"] = "L, U, sd and W are also given as explicit columns; all are integer ticks with U = 20 for the grid protocols."
g = ["Cell value", "Appears for", "Meaning", "L", "U", "sd", "W"]
ws5.append([]); ws5.append(g); style_header(ws5, 4, len(g))
GUIDE = [
 ("regime A", "leader election (plain + inbox)", "Tight intra-datacentre link: narrow transit window, short storage, generous deadline. The regime where timing prunes most.", 18, 20, 2, 200),
 ("regime B", "leader election (plain + inbox)", "Intra-datacentre TCP: moderate transit floor, no storage, modest deadline.", 6, 20, 0, 60),
 ("regime C", "leader election (plain + inbox)", "Inter-availability-zone: low transit floor with a longer tail, mid storage.", 2, 20, 5, 100),
 ("regime F", "leader election (plain + inbox)", "Adversarial / backpressure: widest message window and tightest deadline ratio. COMPLETENESS CONTROL: timed counts must equal the untimed ones exactly.", 0, 20, 10, 40),
 ("B=2 / B=3 regime X", "leader election (plain + inbox)", "Two or three ballots at regime X; the ballot count is in the Rounds/Ballots column.", "", 20, "", ""),
 ("W/U=2", "three-phase commit", "Skeen's minimum bounded wait (W = 2U). Least pruning.", 0, 20, 5, 40),
 ("W/U=3, 5, 7, 10", "three-phase commit", "Deadline 3, 5, 7 or 10 times the worst-case transit; W/U=7 is the default point.", 0, 20, 5, ""),
 ("L/U=0.25", "three-phase commit", "W/U=7 with the transit floor raised to a quarter of U.", 5, 20, 5, 140),
 ("sd/U=0.5", "three-phase commit", "W/U=7 with storage duration doubled to half of U.", 0, 20, 10, 140),
 ("R=2, R=3, R=4", "three-phase commit", "Two, three or four commit rounds at the default point (W/U=7).", 0, 20, 5, 140),
 ("C=.. U=.. L=.. sd=.. TTL=.. pause=.. fence=..", "lease", "C clients, lease TTL, holder pause and fencing on/off; boundary pause = TTL - 2(U - L + sd).", "", "", "", ""),
 ("N=.. U=.. L=.. sd=.. Wp=.. Ws=..", "SWIM", "N members, probe window Wp and suspicion window Ws; boundary Ws = 2U + 2sd - min(L, max(0, 2L - Wp)).", "", "", "", ""),
]
for row in GUIDE:
    ws5.append(list(row)); i = ws5.max_row; border_row(ws5, i, len(g))
    for c in range(1, len(g) + 1):
        ws5.cell(row=i, column=c).alignment = Alignment(wrap_text=True, vertical="top")
    ws5.cell(row=i, column=1).font = BOLD
ws5.column_dimensions["A"].width = 30; ws5.column_dimensions["B"].width = 30; ws5.column_dimensions["C"].width = 90
for col in ("D", "E", "F", "G"):
    ws5.column_dimensions[col].width = 6

out = os.path.join(HERE, "results_v6.xlsx")
wb.save(out)
print(f"wrote {out}")
print(f"  Results:    {len(gcells)} cells ({pending} with a side still running)")
print(f"  Comparison: {len(cmp_cells)} cells")
print(f"  Boundaries: {len(bnd)} rows, {agree} matching")
